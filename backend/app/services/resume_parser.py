import io
import docx
import pypdf
import pypdf.errors

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB


def parse_resume_file(file_name: str, content: bytes) -> str:
    """
    Validate file properties and extract raw text from PDF or DOCX content.

    Raises ValueError with user-friendly message if validation or text extraction fails.
    """
    # 1. File size validation
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ValueError("File size exceeds maximum allowed limit of 5 MB.")

    if len(content) == 0:
        raise ValueError("Uploaded file is empty.")

    # 2. Extension validation
    clean_filename = file_name.lower().strip()
    if clean_filename.endswith(".pdf"):
        return _extract_text_from_pdf(content)
    elif clean_filename.endswith(".docx"):
        return _extract_text_from_docx(content)
    else:
        raise ValueError("Unsupported file format. Only PDF and DOCX files are allowed.")


def _extract_text_from_pdf(content: bytes) -> str:
    try:
        pdf_stream = io.BytesIO(content)
        reader = pypdf.PdfReader(pdf_stream)

        # Check encryption
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
                if not decrypted:
                    raise ValueError("PDF file is encrypted or password-protected.")
            except Exception:
                raise ValueError("PDF file is encrypted or password-protected.")

        extracted_chunks = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                extracted_chunks.append(text.strip())

        full_text = "\n\n".join(extracted_chunks).strip()
        if not full_text:
            raise ValueError(
                "Could not extract text from PDF. The document may be empty, image-only/scanned, or unreadable."
            )

        return full_text

    except ValueError:
        raise
    except pypdf.errors.PyPdfError as e:
        raise ValueError(f"Malformed or corrupt PDF document: {str(e)}")
    except Exception as e:
        raise ValueError(f"Failed to parse PDF document: {str(e)}")


def _extract_text_from_docx(content: bytes) -> str:
    try:
        docx_stream = io.BytesIO(content)
        doc = docx.Document(docx_stream)

        extracted_chunks = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                extracted_chunks.append(text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    extracted_chunks.append(" | ".join(row_texts))

        full_text = "\n\n".join(extracted_chunks).strip()
        if not full_text:
            raise ValueError(
                "Could not extract text from DOCX document. The document contains no readable text."
            )

        return full_text

    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Malformed or unreadable DOCX document: {str(e)}")
